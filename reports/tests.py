from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import Employee, Funkce, Oddeleni, Odbor, Sekce, TypUvazku
from leaves.models import TypStavu, ZadostOStav
from timetracking.models import WorkSession, WorkdaySummary
from reports.services import NEPRITOMEN, PRITOMEN, stav_zamestnance

User = get_user_model()


class PrehledPritomnostiTestCase(TestCase):
    def setUp(self):
        sekce = Sekce.objects.create(nazev="Sekce", kod="S1")
        self.odbor = Odbor.objects.create(sekce=sekce, nazev="Odbor", kod="O1")
        self.oddeleni_a = Oddeleni.objects.create(odbor=self.odbor, nazev="Oddeleni A", kod="OA")
        self.oddeleni_b = Oddeleni.objects.create(odbor=self.odbor, nazev="Oddeleni B", kod="OB")
        self.uvazek = TypUvazku.objects.create(
            nazev="Plny uvazek", hodiny_denne=Decimal("8.00"), hodiny_tyydne=Decimal("40.00")
        )

        self.zam_a = self._vytvor_zamestnance("a@example.com", "Alena", "Adamova", self.oddeleni_a, "1")
        self.zam_b = self._vytvor_zamestnance("b@example.com", "Bedrich", "Bily", self.oddeleni_b, "2")

        self.zam_a.funkce = Funkce.objects.get(kod=Funkce.VEDOUCI_ODDELENI)
        self.zam_a.save()

        self.typ_dovolena = TypStavu.objects.create(
            nazev="Dovolena", zkratka="DOV", odecita_ze_zustatku=True,
            kategorie_pro_prehled=TypStavu.KategoriePrehled.DOVOLENA,
            barva="#FFC107",
        )
        self.typ_home_office = TypStavu.objects.create(
            nazev="Home office", zkratka="HO",
            odecita_ze_zustatku=False, je_pritomnost=True,
            vyzaduje_schvaleni=False, barva="#6610F2",
        )

    def _vytvor_zamestnance(self, email, first, last, oddeleni, cislo):
        user = User.objects.create_user(
            username=email, email=email, first_name=first, last_name=last,
        )
        return Employee.objects.create(
            user=user, osobni_cislo=cislo, oddeleni=oddeleni,
            typ_uvazku=self.uvazek, datum_nastupu=date(2020, 1, 1),
        )

    def _prihlas(self, employee, email):
        employee.user.set_password("test12345")
        employee.user.save()
        self.assertTrue(self.client.login(username=email, password="test12345"))

    def test_pritomnost_ma_prednost_pred_schvalenou_absenci(self):
        dnes = timezone.localdate()
        ZadostOStav.objects.create(
            employee=self.zam_b, typ=self.typ_dovolena,
            datum_od=dnes, datum_do=dnes, stav=ZadostOStav.Stav.SCHVALENO,
        )
        WorkSession.objects.create(
            employee=self.zam_b,
            zacatek=timezone.now().replace(hour=8, minute=0, second=0, microsecond=0),
        )
        stav = stav_zamestnance(self.zam_b, dnes)
        self.assertEqual(stav.kod, PRITOMEN)

    def test_bez_pritomnosti_a_bez_zadosti_je_nepritomen(self):
        stav = stav_zamestnance(self.zam_b, timezone.localdate())
        self.assertEqual(stav.kod, NEPRITOMEN)

    def test_schvalena_zadost_bez_pritomnosti_ukaze_typ_stavu(self):
        dnes = timezone.localdate()
        ZadostOStav.objects.create(
            employee=self.zam_b, typ=self.typ_dovolena,
            datum_od=dnes, datum_do=dnes, stav=ZadostOStav.Stav.SCHVALENO,
        )
        stav = stav_zamestnance(self.zam_b, dnes)
        self.assertEqual(stav.kod, self.typ_dovolena.zkratka)
        self.assertEqual(stav.popisek, self.typ_dovolena.nazev)
        self.assertEqual(stav.barva, self.typ_dovolena.barva)

    def test_home_office_ma_prednost_pred_pritomen(self):
        """Home office (je_pritomnost=True) se zobrazí i když existuje WorkSession."""
        dnes = timezone.localdate()
        ZadostOStav.objects.create(
            employee=self.zam_b, typ=self.typ_home_office,
            datum_od=dnes, datum_do=dnes, stav=ZadostOStav.Stav.SCHVALENO,
        )
        WorkSession.objects.create(
            employee=self.zam_b,
            zacatek=timezone.now().replace(hour=8, minute=0, second=0, microsecond=0),
        )
        stav = stav_zamestnance(self.zam_b, dnes)
        self.assertEqual(stav.kod, self.typ_home_office.zkratka)

    def test_home_office_se_zobrazi_i_bez_worksession(self):
        dnes = timezone.localdate()
        ZadostOStav.objects.create(
            employee=self.zam_b, typ=self.typ_home_office,
            datum_od=dnes, datum_do=dnes, stav=ZadostOStav.Stav.SCHVALENO,
        )
        stav = stav_zamestnance(self.zam_b, dnes)
        self.assertEqual(stav.kod, self.typ_home_office.zkratka)

    def test_novy_typ_zalozeny_jen_v_testu_funguje_bez_zasahu_do_kodu(self):
        """Nový TypStavu (bez odkazu na jakýkoli enum) se v přehledu zobrazí správně."""
        dnes = timezone.localdate()
        typ_sjezd = TypStavu.objects.create(
            nazev="Sjezd", zkratka="SJ",
            odecita_ze_zustatku=False, je_pritomnost=False,
            vyzaduje_schvaleni=False, barva="#123456",
        )
        ZadostOStav.objects.create(
            employee=self.zam_b, typ=typ_sjezd,
            datum_od=dnes, datum_do=dnes, stav=ZadostOStav.Stav.SCHVALENO,
        )
        stav = stav_zamestnance(self.zam_b, dnes)
        self.assertEqual(stav.kod, "SJ")
        self.assertEqual(stav.popisek, "Sjezd")
        self.assertEqual(stav.barva, "#123456")

    @staticmethod
    def _videni_zamestnanci(response):
        """Zploští non-staff `skupiny` ({'oddeleni', 'radky'}) na set ID zaměstnanců."""
        videni = set()
        for skupina in response.context["skupiny"]:
            for radek in skupina["radky"]:
                videni.add(radek["employee"].pk)
        return videni

    def test_radovy_zamestnanec_vidi_cely_svuj_odbor(self):
        self._prihlas(self.zam_b, "b@example.com")

        response = self.client.get(reverse("reports:prehled_pritomnosti"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._videni_zamestnanci(response), {self.zam_a.pk, self.zam_b.pk})

    def test_vedouci_oddeleni_vidi_cely_odbor_kdyz_je_priznak_zapnuty(self):
        """Zapnutý Odbor.zamestnanci_vidi_cely_odbor (výchozí) platí i pro vedoucí oddělení —
        v read-only přehledech nejsou omezeni jen na svou CRUD správu (vlastní oddělení)."""
        self._prihlas(self.zam_a, "a@example.com")

        response = self.client.get(reverse("reports:prehled_pritomnosti"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._videni_zamestnanci(response), {self.zam_a.pk, self.zam_b.pk})

    def test_vedouci_oddeleni_vidi_jen_sve_oddeleni_kdyz_je_priznak_vypnuty(self):
        self.odbor.zamestnanci_vidi_cely_odbor = False
        self.odbor.save()
        self._prihlas(self.zam_a, "a@example.com")

        response = self.client.get(reverse("reports:prehled_pritomnosti"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._videni_zamestnanci(response), {self.zam_a.pk})

    def test_legenda_ma_stejnou_barvu_jako_stav_zamestnance(self):
        dnes = timezone.localdate()
        ZadostOStav.objects.create(
            employee=self.zam_b, typ=self.typ_dovolena,
            datum_od=dnes, datum_do=dnes, stav=ZadostOStav.Stav.SCHVALENO,
        )
        self._prihlas(self.zam_a, "a@example.com")

        response = self.client.get(reverse("reports:prehled_pritomnosti"))
        self.assertEqual(response.status_code, 200)
        legenda = {p["stav"].kod: p["stav"] for p in response.context["pocty"]}
        self.assertEqual(legenda[self.typ_dovolena.zkratka].barva, self.typ_dovolena.barva)
        # badge s barvou dovolené se vykreslí 2× — v legendě a u řádku zaměstnance
        self.assertContains(response, f"background-color: {self.typ_dovolena.barva};", count=2)

    def test_prehled_tymu_sdili_stejny_rozsah_jako_pritomnost(self):
        """Issue #28 — Odbor (prehled_tymu) musí ukazovat stejný okruh lidí jako
        Přítomnost (viditelni_zamestnanci), ne jen CRUD-spravované podřízené."""
        self._prihlas(self.zam_b, "b@example.com")

        response = self.client.get(reverse("reports:prehled_tymu"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._videni_zamestnanci(response), {self.zam_a.pk, self.zam_b.pk})

    def test_vyhledani_najde_zamestnance_mimo_vlastni_odbor(self):
        jina_sekce = Sekce.objects.create(nazev="Jina sekce", kod="S2")
        jiny_odbor = Odbor.objects.create(sekce=jina_sekce, nazev="Jiny odbor", kod="O2")
        jine_oddeleni = Oddeleni.objects.create(odbor=jiny_odbor, nazev="Jine oddeleni", kod="OC")
        zam_c = self._vytvor_zamestnance("c@example.com", "Cyril", "Cerny", jine_oddeleni, "3")

        self._prihlas(self.zam_a, "a@example.com")

        response = self.client.get(reverse("reports:vyhledat_zamestnance"), {"q": "Cerny"})
        self.assertEqual(response.status_code, 200)
        vysledky = {v["employee"].pk for v in response.context["vysledky"]}
        self.assertEqual(vysledky, {zam_c.pk})

    def test_prazdny_dotaz_nevraci_zadne_vysledky(self):
        self._prihlas(self.zam_a, "a@example.com")
        response = self.client.get(reverse("reports:vyhledat_zamestnance"))
        self.assertEqual(response.context["vysledky"], [])


class PrehledTymuAExportBilanceTestCase(TestCase):
    """Issue #34 — Odbor a XLSX export ukazují přesčas/nedostatek zvlášť, bez záporných hodin."""

    def setUp(self):
        sekce = Sekce.objects.create(nazev="Sekce", kod="S3")
        odbor = Odbor.objects.create(sekce=sekce, nazev="Odbor", kod="O3")
        oddeleni = Oddeleni.objects.create(odbor=odbor, nazev="Oddeleni", kod="OD3")
        uvazek = TypUvazku.objects.create(
            nazev="Plny uvazek", hodiny_denne=Decimal("8.00"), hodiny_tyydne=Decimal("40.00")
        )
        user = User.objects.create_user(
            username="d@example.com", email="d@example.com",
            first_name="Dana", last_name="Dvorakova", password="test12345",
        )
        self.zam = Employee.objects.create(
            user=user, osobni_cislo="4", oddeleni=oddeleni,
            typ_uvazku=uvazek, datum_nastupu=date(2020, 1, 1),
        )
        # Po 7.9. (+45), Út 8.9. (-90), rozpracovaný den 9.9. (bez uzavřeného bloku)
        WorkdaySummary.objects.create(
            employee=self.zam, datum=date(2026, 9, 7),
            hrube_minuty=525, odpracovane_minuty=525, prescos_minuty=45,
        )
        WorkdaySummary.objects.create(
            employee=self.zam, datum=date(2026, 9, 8),
            hrube_minuty=390, odpracovane_minuty=390, prescos_minuty=-90,
        )
        WorkdaySummary.objects.create(
            employee=self.zam, datum=date(2026, 9, 9),
            hrube_minuty=0, odpracovane_minuty=0, prescos_minuty=-480,
        )
        self.client = Client()
        user.set_password("test12345")
        user.save()
        self.assertTrue(self.client.login(username="d@example.com", password="test12345"))

    def test_odbor_ukazuje_prescas_a_nedostatek_zvlast_bez_zapornych_hodin(self):
        response = self.client.get(reverse("reports:prehled_tymu"), {"rok": 2026, "mesic": 9})
        self.assertEqual(response.status_code, 200)
        obsah = response.content.decode("utf-8")
        self.assertNotIn("-2h", obsah)
        self.assertNotIn("-1h", obsah)
        self.assertIn("Nedostatek", obsah)

        (skupina,) = response.context["skupiny"]
        (radek,) = skupina["radky"]
        self.assertEqual((radek["prescas_minuty"], radek["nedostatek_minuty"]), (45, 90))

    def test_export_xlsx_ma_tydenni_a_mesicni_souhrn_bez_zapornych_hodin(self):
        import openpyxl
        from io import BytesIO

        response = self.client.get(reverse("reports:export_xlsx"), {"rok": 2026, "mesic": 9})
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        hlavicka = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(hlavicka, ["Datum", "Den", "Odpracováno", "Přesčas", "Nedostatek", "Svátek/Víkend", "Bilance"])

        popisky = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value]
        tydenni_radky = [p for p in popisky if p and p.startswith("Týden")]
        self.assertEqual(len(tydenni_radky), 1)
        self.assertIn("Celkem za měsíc", popisky)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    self.assertNotIn("-2h", cell.value)

        radek_mesic = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "Celkem za měsíc")
        self.assertEqual(radek_mesic[3].value, "0h 45min")   # Přesčas
        self.assertEqual(radek_mesic[4].value, "1h 30min")   # Nedostatek
        self.assertEqual(radek_mesic[6].value, "−0h 45min")  # Bilance
